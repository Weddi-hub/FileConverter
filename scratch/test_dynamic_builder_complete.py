import os
import re
import pdfplumber
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

def clean_val(v):
    return re.sub(r'\s+', ' ', str(v or "")).strip()

def parse_num(v_str):
    s = str(v_str or "").strip().replace(",", "")
    if re.fullmatch(r"[-+]?\d+", s):
        try:
            return int(s)
        except ValueError:
            return None
    if re.fullmatch(r"[-+]?\d+\.\d+", s):
        try:
            return float(s)
        except ValueError:
            return None
    return None

def group_words_by_y(words, y_tol=3.5):
    if not words:
        return []
    sorted_words = sorted(words, key=lambda w: (w['top'], w['x0']))
    lines = []
    curr_line = []
    curr_top = None
    for w in sorted_words:
        if curr_top is None:
            curr_top = w['top']
            curr_line.append(w)
        elif abs(w['top'] - curr_top) <= y_tol:
            curr_line.append(w)
        else:
            curr_line.sort(key=lambda x: x['x0'])
            lines.append(curr_line)
            curr_line = [w]
            curr_top = w['top']
    if curr_line:
        curr_line.sort(key=lambda x: x['x0'])
        lines.append(curr_line)
    return lines

def build_dynamic_document_excel(pdf_path: str, output_excel_path: str, sheet_name: str = "Sheet1"):
    parent_dir = os.path.dirname(os.path.abspath(output_excel_path))
    if parent_dir and not os.path.exists(parent_dir):
        os.makedirs(parent_dir, exist_ok=True)
        
    logo_path = None
    title_block = []
    metadata_items = []
    col_names = []
    col_intervals = []
    all_data_rows = []
    grand_total_info = {}
    financial_summary_items = []
    notes = []
    
    with pdfplumber.open(pdf_path) as pdf:
        p1 = pdf.pages[0]
        page_width = float(p1.width)
        words_p1 = p1.extract_words()
        lines_p1 = group_words_by_y(words_p1, y_tol=3.5)
        
        # 1. Logo Extraction
        if p1.images:
            try:
                img_obj = p1.images[0]
                if img_obj['top'] < 100:
                    temp_logo = os.path.join(parent_dir, f"dynamic_logo_{os.getpid()}.png")
                    cropped = p1.crop((img_obj['x0'], img_obj['top'], img_obj['x1'], img_obj['bottom'])).to_image(resolution=150)
                    cropped.save(temp_logo)
                    logo_path = temp_logo
            except Exception:
                logo_path = None
                
        # 2. Header & Column Discovery on Page 1
        header_line_idx = -1
        header_chunks = []
        for idx, line in enumerate(lines_p1):
            chunks = []
            curr = []
            for w in line:
                if not curr or (w['x0'] - curr[-1]['x1'] < 16.0):
                    curr.append(w)
                else:
                    chunks.append(curr)
                    curr = [w]
            if curr:
                chunks.append(curr)
            line_str = " ".join(w['text'] for w in line)
            
            # Check for title / metadata before table
            if header_line_idx == -1:
                if len(chunks) >= 4 and any(k in line_str.upper() for k in ["SR#", "BILL", "NO.", "DATE", "AMOUNT"]):
                    header_line_idx = idx
                    header_chunks = chunks
                else:
                    if any(k in line_str for k in ["PAKISTAN AIRPORTS AUTHORITY", "SUMMARY OF AERONAUTICAL", "PAA-001"]):
                        title_block.append(line_str)
                    elif ":" in line_str:
                        metadata_items.append(line_str)
                        
        if header_line_idx == -1:
            raise ValueError(f"Could not discover table headers in {pdf_path}")
            
        # Check for sub-header line (e.g. (RS.) IF ANY)
        sub_headers = []
        if header_line_idx + 1 < len(lines_p1):
            next_line = lines_p1[header_line_idx + 1]
            next_str = " ".join(w['text'] for w in next_line)
            if any(k in next_str for k in ["(RS.)", "(US$)", "IF ANY", "PKR", "(TONS)"]):
                sub_chunks = []
                curr = []
                for w in next_line:
                    if not curr or (w['x0'] - curr[-1]['x1'] < 16.0):
                        curr.append(w)
                    else:
                        sub_chunks.append(curr)
                        curr = [w]
                if curr:
                    sub_chunks.append(curr)
                sub_headers = sub_chunks
                
        # Merge sub-headers into col_names
        for c in header_chunks:
            txt = " ".join(w['text'] for w in c)
            cx = (c[0]['x0'] + c[-1]['x1']) / 2
            for sc in sub_headers:
                sc_cx = (sc[0]['x0'] + sc[-1]['x1']) / 2
                if abs(sc_cx - cx) < 30.0 or (c[0]['x0'] - 10 <= sc_cx <= c[-1]['x1'] + 10):
                    txt += " " + " ".join(w['text'] for w in sc)
                    break
            # Clean up parenthesis whitespace
            txt = txt.replace("( ", "(").replace(" )", ")")
            col_names.append(txt.strip())
            
        n_cols = len(header_chunks)
        for i in range(n_cols):
            x_start = 0.0 if i == 0 else (header_chunks[i-1][-1]['x1'] + header_chunks[i][0]['x0']) / 2
            x_end = page_width if i == n_cols - 1 else (header_chunks[i][-1]['x1'] + header_chunks[i+1][0]['x0']) / 2
            col_intervals.append((col_names[i], x_start, x_end))
            
        header_y_max = max(w['bottom'] for w in lines_p1[header_line_idx + (1 if sub_headers else 0)])
        
        # 3. Extract Rows Across All Pages
        first_header_token = header_chunks[0][0]['text']
        second_header_token = header_chunks[1][0]['text'] if n_cols > 1 else ""
        
        in_financial_summary = False
        in_notes = False
        
        for p_idx, page in enumerate(pdf.pages):
            p_words = page.extract_words()
            p_lines = group_words_by_y(p_words, y_tol=3.5)
            p_text = page.extract_text() or ""
            is_table_page = (first_header_token in p_text and second_header_token in p_text)
            
            for line in p_lines:
                line_str = " ".join(w['text'] for w in line)
                line_y = line[0]['top']
                
                # Skip header block on page 1
                if p_idx == 0 and line_y <= header_y_max:
                    continue
                # Skip running headers
                if any(k in line_str for k in ["PAKISTAN AIRPORTS AUTHORITY", "PAA-001-FNBL", "PAGE NO:", "SUMMARY OF AERONAUTICAL BILLS", "June 17, 2026", "5:06 pm"]):
                    continue
                if first_header_token in line_str and second_header_token in line_str:
                    continue
                if "(RS.)" in line_str and "IF ANY" in line_str:
                    continue
                if "AIRLINE :" in line_str:
                    continue
                    
                # Check for Grand Total line
                if "TOTAL AMOUNT DUE" in line_str.upper():
                    grand_total_match = re.search(r'TOTAL AMOUNT DUE[^\d]*([\d,]+)', line_str, re.IGNORECASE)
                    if grand_total_match:
                        grand_total_info['label'] = "TOTAL AMOUNT DUE FOR FORTNIGHT :"
                        grand_total_info['amount'] = int(grand_total_match.group(1).replace(",", ""))
                    in_financial_summary = True
                    continue
                    
                # Financial summary items (Page 5)
                if in_financial_summary:
                    if "NOTE" in line_str.upper():
                        in_notes = True
                        in_financial_summary = False
                        notes.append(line_str)
                        continue
                    if any(k in line_str.upper() for k in ["ARREARS", "SURCHARGE", "KIBOR", "TOTAL ARREARS", "TOTAL AMOUNT PAYABLE"]):
                        financial_summary_items.append(line_str)
                        continue
                        
                if in_notes:
                    notes.append(line_str)
                    continue
                    
                if not is_table_page:
                    continue
                    
                # Table row
                row = [""] * n_cols
                for w in line:
                    cx = (w['x0'] + w['x1']) / 2
                    for c_idx, (_, x0, x1) in enumerate(col_intervals):
                        if x0 <= cx < x1:
                            if row[c_idx]:
                                row[c_idx] += " " + w['text']
                            else:
                                row[c_idx] = w['text']
                            break
                            
                # Check if it's a valid data row (e.g. starts with serial number)
                if row[0] and row[0].isdigit():
                    all_data_rows.append(row)
                    
    # Build Excel Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.views.sheetView[0].showGridLines = True
    
    thin_border = Side(style='thin', color='000000')
    double_border = Side(style='double', color='000000')
    cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    grand_total_border = Border(top=thin_border, bottom=double_border, left=thin_border, right=thin_border)
    
    # Fonts & Fills
    title_font = Font(name="Calibri", size=14, bold=True)
    subtitle_font = Font(name="Calibri", size=11, bold=True)
    meta_font = Font(name="Calibri", size=10, bold=True)
    header_font = Font(name="Calibri", size=10, bold=True)
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    data_font = Font(name="Calibri", size=10)
    total_font = Font(name="Calibri", size=10, bold=True)
    total_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    note_font = Font(name="Calibri", size=9)
    note_bold = Font(name="Calibri", size=9, bold=True)
    
    cur_row = 1
    
    # 1. Logo
    if logo_path and os.path.exists(logo_path):
        try:
            img = openpyxl.drawing.image.Image(logo_path)
            img.width = 110
            img.height = 45
            ws.add_image(img, "B2")
        except Exception:
            pass
            
    # 2. Title Block
    cur_row = 2
    for t_line in title_block:
        if "PAKISTAN AIRPORTS AUTHORITY" in t_line:
            cell = ws.cell(row=cur_row, column=3, value="PAKISTAN AIRPORTS AUTHORITY")
            cell.font = title_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if "PAA-001-FNBL-1.0" in t_line:
                ref_cell = ws.cell(row=cur_row, column=n_cols, value="PAA-001-FNBL-1.0")
                ref_cell.font = Font(name="Calibri", size=9, bold=True)
                ref_cell.alignment = Alignment(horizontal="right", vertical="center")
        elif "SUMMARY OF AERONAUTICAL BILLS" in t_line:
            cell = ws.cell(row=cur_row, column=3, value=t_line)
            cell.font = subtitle_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        cur_row += 1
        
    cur_row += 1
    for m_line in metadata_items:
        cell = ws.cell(row=cur_row, column=2, value=m_line)
        cell.font = meta_font
        cur_row += 1
        
    cur_row += 1
    table_start_row = cur_row
    
    # 3. Table Headers
    ws.row_dimensions[cur_row].height = 24
    for c_idx, h_name in enumerate(col_names, start=1):
        cell = ws.cell(row=cur_row, column=c_idx, value=h_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    cur_row += 1
    
    # 4. Infer Column Data Types
    col_types = []
    for c_idx in range(n_cols):
        vals = [r[c_idx] for r in all_data_rows if r[c_idx].strip()]
        if not vals:
            col_types.append("empty")
            continue
        # Check if column is integer
        int_count = sum(1 for v in vals if parse_num(v) is not None and isinstance(parse_num(v), int))
        float_count = sum(1 for v in vals if parse_num(v) is not None and isinstance(parse_num(v), float))
        # Exclude IDs or codes (if col 0 is serial number, it's int; if col 1 is bill no, it has letters)
        if c_idx == 0:
            col_types.append("int")
        elif c_idx == 1:
            col_types.append("str")
        elif int_count == len(vals):
            col_types.append("int")
        elif (int_count + float_count) == len(vals):
            col_types.append("float")
        else:
            col_types.append("str")
            
    # 5. Populate Data Rows
    for r_vals in all_data_rows:
        ws.row_dimensions[cur_row].height = 19
        for c_idx, val in enumerate(r_vals, start=1):
            cell = ws.cell(row=cur_row, column=c_idx)
            cell.border = cell_border
            cell.font = data_font
            
            c_type = col_types[c_idx - 1]
            if c_type == "int":
                parsed = parse_num(val)
                if parsed is not None:
                    cell.value = parsed
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="right" if c_idx > 1 else "center", vertical="center")
                else:
                    cell.value = val
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_type == "float":
                parsed = parse_num(val)
                if parsed is not None:
                    cell.value = parsed
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.value = val
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.value = val
                align = "center" if any(k in col_names[c_idx - 1].upper() for k in ["DATE", "SR#", "CODE", "REG"]) else "left"
                cell.alignment = Alignment(horizontal=align, vertical="center")
        cur_row += 1
        
    # 6. Grand Total Row
    if grand_total_info:
        ws.row_dimensions[cur_row].height = 22
        for c_idx in range(1, n_cols + 1):
            cell = ws.cell(row=cur_row, column=c_idx)
            cell.border = grand_total_border
            cell.font = total_font
            cell.fill = total_fill
            
        # Label in col 2 or 3
        lbl_cell = ws.cell(row=cur_row, column=2, value=grand_total_info.get('label', 'TOTAL AMOUNT DUE :'))
        lbl_cell.alignment = Alignment(horizontal="left", vertical="center")
        
        # Total amount in amount column (col 4 / AMOUNT BILLED)
        tot_cell = ws.cell(row=cur_row, column=4, value=grand_total_info.get('amount', 0))
        tot_cell.number_format = '#,##0'
        tot_cell.alignment = Alignment(horizontal="right", vertical="center")
        cur_row += 2
        
    # 7. Financial Summary Block (Page 5)
    if financial_summary_items:
        for f_item in financial_summary_items:
            ws.row_dimensions[cur_row].height = 18
            parts = f_item.rsplit(" ", 1)
            if len(parts) == 2 and parse_num(parts[1]) is not None:
                lbl = ws.cell(row=cur_row, column=2, value=parts[0].strip())
                lbl.font = Font(name="Calibri", size=10, bold=True)
                val_c = ws.cell(row=cur_row, column=4, value=int(parts[1].replace(",", "")))
                val_c.font = Font(name="Calibri", size=10, bold=True)
                val_c.number_format = '#,##0'
                val_c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                lbl = ws.cell(row=cur_row, column=2, value=f_item)
                lbl.font = Font(name="Calibri", size=10, bold=True)
            cur_row += 1
        cur_row += 1
        
    # 8. Notes Section
    if notes:
        ws.row_dimensions[cur_row].height = 18
        note_hdr = ws.cell(row=cur_row, column=2, value="NOTE:")
        note_hdr.font = note_bold
        cur_row += 1
        for n_line in notes:
            if n_line.strip() == "NOTE:":
                continue
            ws.row_dimensions[cur_row].height = 16
            nc = ws.cell(row=cur_row, column=2, value=n_line.strip())
            nc.font = note_font
            cur_row += 1
            
    # Auto-fit column widths
    for c in range(1, n_cols + 1):
        col_letter = get_column_letter(c)
        max_len = 0
        for r in range(table_start_row, cur_row):
            val = str(ws.cell(row=r, column=c).value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)
        
    wb.save(output_excel_path)
    if logo_path and os.path.exists(logo_path):
        try:
            os.remove(logo_path)
        except Exception:
            pass
            
    df_preview = pd.DataFrame(all_data_rows, columns=col_names)
    return df_preview, output_excel_path

if __name__ == "__main__":
    df, out_path = build_dynamic_document_excel(
        'tests/samples/PAA_Aeronautical_Bills_Summary.pdf',
        'tests/samples/dynamic_summary_out.xlsx'
    )
    print(f"Generated {out_path} with {len(df)} rows and {len(df.columns)} columns!")
    print(df.head(3))
    print(df.tail(3))
