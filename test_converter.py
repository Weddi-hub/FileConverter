"""
Comprehensive Test Suite for pdf_to_excel Converter
===================================================
Tests universal adaptive table extraction, borderless parsing,
multi-page continuation, single continuous sheet constraint,
index omission, specialized PAA bills, and storage lifecycle.
"""

import os
import tempfile
import unittest
import openpyxl
import pandas as pd
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

from pdf_to_excel import convert_pdf_to_excel, clean_cell_text, clean_table
from app import purge_expired_files, UPLOAD_DIR, OUTPUT_DIR


class TestPDFToExcelConverter(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root_dir = os.path.dirname(os.path.abspath(__file__))
        cls.samples_dir = os.path.join(cls.root_dir, "tests", "samples")

        # Create a temporary directory for any generated test files so root stays 100% clean
        cls.temp_dir_obj = tempfile.TemporaryDirectory()
        cls.temp_dir = cls.temp_dir_obj.name

        cls.empty_pdf = os.path.join(cls.temp_dir, "test_empty.pdf")
        styles = getSampleStyleSheet()
        doc_e = SimpleDocTemplate(cls.empty_pdf, pagesize=letter)
        doc_e.build([Paragraph("This is a document with no tables anywhere.", styles['Normal'])])

    @classmethod
    def tearDownClass(cls):
        cls.temp_dir_obj.cleanup()

    def test_adaptive_borderless_extraction(self):
        """Verify borderless table extraction using Adaptive Spatial Phrase Clusterer."""
        pdf_path = os.path.join(self.samples_dir, "borderless_sample.pdf")
        if not os.path.exists(pdf_path):
            self.skipTest("borderless_sample.pdf not found in tests/samples")

        out_xlsx = os.path.join(self.temp_dir, "borderless_out.xlsx")
        df, path = convert_pdf_to_excel(pdf_path, output_excel_path=out_xlsx)

        self.assertTrue(os.path.exists(out_xlsx))
        wb = openpyxl.load_workbook(out_xlsx)
        self.assertEqual(wb.sheetnames, ["Sheet1"], "Must be on a single continuous sheet 'Sheet1'")
        wb.close()

        excel_df = pd.read_excel(out_xlsx, sheet_name="Sheet1")
        expected_cols = ["Employee ID", "Full Name", "Role", "Salary"]
        self.assertEqual(list(excel_df.columns), expected_cols)
        self.assertEqual(len(excel_df), 3)
        self.assertEqual(list(excel_df["Employee ID"]), ["E001", "E002", "E003"])

    def test_complex_multipage_and_mixed_tables(self):
        """Verify multi-page table continuation and distinct tables on page 5."""
        pdf_path = os.path.join(self.samples_dir, "complex_sample.pdf")
        if not os.path.exists(pdf_path):
            self.skipTest("complex_sample.pdf not found in tests/samples")

        out_xlsx = os.path.join(self.temp_dir, "complex_out.xlsx")
        df, path = convert_pdf_to_excel(pdf_path, output_excel_path=out_xlsx)

        self.assertTrue(os.path.exists(out_xlsx))
        wb = openpyxl.load_workbook(out_xlsx)
        self.assertEqual(wb.sheetnames, ["Sheet1"])
        ws = wb["Sheet1"]

        # Row 1 must have Table 1 headers
        self.assertEqual([ws.cell(1, c).value for c in range(1, 5)], ["Region", "Product", "Q1 Sales", "Q2 Sales"])
        # Continuation rows across pages 1, 2, 3
        self.assertEqual(ws.cell(2, 1).value, "North")
        self.assertEqual(ws.cell(4, 1).value, "East")
        self.assertEqual(ws.cell(6, 1).value, "Central")
        # Row 9 has Table 2 headers (from page 5)
        self.assertEqual([ws.cell(9, c).value for c in range(1, 4)], ["Dept ID", "Department Name", "Headcount"])
        wb.close()

    def test_paa_air_navigation_bill(self):
        """Verify specialized conversion for 8-page PAA Air Navigation Bill."""
        pdf_path = os.path.join(self.samples_dir, "8pages_air_navigation.pdf")
        if not os.path.exists(pdf_path):
            self.skipTest("8pages_air_navigation.pdf not found in tests/samples")

        out_xlsx = os.path.join(self.temp_dir, "air_nav_out.xlsx")
        df, path = convert_pdf_to_excel(pdf_path, output_excel_path=out_xlsx)

        self.assertTrue(os.path.exists(out_xlsx))
        wb = openpyxl.load_workbook(out_xlsx)
        self.assertEqual(wb.sheetnames, ["Sheet1"])
        ws = wb["Sheet1"]
        self.assertEqual(ws["C1"].value, "PAKISTAN AIRPORTS AUTHORITY")
        self.assertEqual(ws["C2"].value, "AIR NAVIGATION FOR LANDING")
        self.assertEqual(ws["A11"].value, "BILL ITEM ID")
        self.assertEqual(ws["L11"].value, "AMOUNT (Rs)")
        self.assertGreater(ws.max_row, 130)

        # Total 140 flights
        flight_rows = [r for r in range(12, ws.max_row + 1) if ws.cell(r, 1).value and str(ws.cell(r, 1).value).isdigit()]
        self.assertEqual(len(flight_rows), 140, "Must extract all 140 flight records in PAA Air Navigation Bill")

        # Verify numeric typing and exact mathematical sum in Rupees
        rs_amounts = [ws.cell(r, 12).value for r in flight_rows]
        self.assertTrue(all(isinstance(v, (int, float)) for v in rs_amounts), "Col 12 must be native numeric")
        self.assertEqual(sum(rs_amounts), 35271810, "Sum of all 140 flights in Rupees must equal exactly 35,271,810")
        wb.close()

    def test_paa_landing_and_housing_bill(self):
        """Verify specialized conversion for PAA Landing & Housing Bill."""
        pdf_path = os.path.join(self.samples_dir, "PAA_Domestic_Bill.pdf")
        if not os.path.exists(pdf_path):
            self.skipTest("PAA_Domestic_Bill.pdf not found in tests/samples")

        out_xlsx = os.path.join(self.temp_dir, "lh_out.xlsx")
        df, path = convert_pdf_to_excel(pdf_path, output_excel_path=out_xlsx)

        self.assertTrue(os.path.exists(out_xlsx))
        wb = openpyxl.load_workbook(out_xlsx)
        self.assertEqual(wb.sheetnames, ["Sheet1"])
        ws = wb["Sheet1"]
        self.assertEqual(ws["C1"].value, "PAKISTAN AIRPORTS AUTHORITY")
        self.assertEqual(ws["C2"].value, "LANDING AND HOUSING")
        self.assertEqual(ws["A10"].value, "BILL ITEM ID")
        self.assertIn("TOTAL CHARGES", str(ws["O10"].value))
        self.assertGreater(ws.max_row, 40)

        # Ensure row 148493004 is completely parsed with all columns
        row_148493004 = None
        for r in range(12, ws.max_row + 1):
            if ws.cell(r, 1).value == "148493004":
                row_148493004 = [ws.cell(r, c).value for c in range(1, 16)]
                break
        self.assertIsNotNone(row_148493004, "Entry 148493004 must be parsed")
        self.assertEqual(row_148493004[0], "148493004")
        self.assertEqual(row_148493004[1], "A-320")
        self.assertEqual(row_148493004[2], "APEDA")
        self.assertIn(row_148493004[3], ["74", 74])
        self.assertEqual(row_148493004[4], "PA703")
        self.assertEqual(row_148493004[5], "12/6/26")
        self.assertEqual(row_148493004[6], "12/6/26")
        self.assertIn(row_148493004[10], ["6,882", 6882])
        self.assertIn(row_148493004[11], ["0", 0])
        self.assertIn(row_148493004[12], ["1,274", 1274])
        self.assertIn(row_148493004[13], ["750", 750])
        self.assertIn(row_148493004[14], ["8,906", 8906])

        # Total 44 flights
        flight_ids = [ws.cell(r, 1).value for r in range(12, ws.max_row + 1) if ws.cell(r, 1).value and str(ws.cell(r, 1).value).isdigit()]
        self.assertEqual(len(flight_ids), 44, "Must extract all 44 flight records in PAA Domestic Bill")
        wb.close()

    def test_dammam_landing_charge_invoice(self):
        """Verify specialized conversion for Dammam Airports Landing Charge Invoice (20 columns)."""
        pdf_path = os.path.join(self.samples_dir, "dammam_landing_charge_invoice.pdf")
        if not os.path.exists(pdf_path):
            self.skipTest("dammam_landing_charge_invoice.pdf not found in tests/samples")

        out_xlsx = os.path.join(self.temp_dir, "dammam_out.xlsx")
        df, path = convert_pdf_to_excel(pdf_path, output_excel_path=out_xlsx)

        self.assertTrue(os.path.exists(out_xlsx))
        wb = openpyxl.load_workbook(out_xlsx)
        self.assertEqual(wb.sheetnames, ["Sheet1"])
        ws = wb["Sheet1"]
        self.assertIn("DAMMAM AIRPORTS", str(ws["C1"].value))
        self.assertIn("LANDING CHARGE INVOICE", str(ws["C2"].value))
        self.assertEqual(len(df), 18)
        self.assertEqual(len(df.columns), 20)
        self.assertEqual(df.columns[0], "Flight Total (SAR)")
        self.assertEqual(df.columns[-1], "Flight No.")
        wb.close()

    def test_jeddah_landing_charge_invoice(self):
        """Verify specialized conversion for Jeddah Airports Landing Charge Invoice (22 columns)."""
        pdf_path = os.path.join(self.samples_dir, "jeddah_landing_charge_invoice.pdf")
        if not os.path.exists(pdf_path):
            self.skipTest("jeddah_landing_charge_invoice.pdf not found in tests/samples")

        out_xlsx = os.path.join(self.temp_dir, "jeddah_out.xlsx")
        df, path = convert_pdf_to_excel(pdf_path, output_excel_path=out_xlsx)

        self.assertTrue(os.path.exists(out_xlsx))
        wb = openpyxl.load_workbook(out_xlsx)
        self.assertEqual(wb.sheetnames, ["Sheet1"], "Must be on a single continuous sheet 'Sheet1'")
        ws = wb["Sheet1"]
        self.assertIn("JEDDAH AIRPORTS", str(ws["C1"].value))
        self.assertIn("LANDING CHARGE INVOICE", str(ws["C2"].value))
        self.assertEqual(len(df), 28, "Must extract all 28 flight records")
        self.assertEqual(len(df.columns), 22, "Must contain all 22 columns including Baggage Amount & Count")
        self.assertEqual(df.columns[0], "Flight Total (SAR)")
        self.assertEqual(df.columns[1], "Baggage Amount")
        self.assertEqual(df.columns[-1], "Flight No.")
        # Verify subtotal and total SAR
        self.assertIn("36,870", str(ws.cell(ws.max_row, 2).value))
    def test_sans_air_navigation_invoice(self):
        """Verify specialized conversion for Saudi Air Navigation Services (SANS) invoice (22 pages, 326 flights, 22 cols)."""
        pdf_path = os.path.join(self.samples_dir, "sans_air_navigation_invoice.pdf")
        if not os.path.exists(pdf_path):
            self.skipTest("sans_air_navigation_invoice.pdf not found in tests/samples")

        out_xlsx = os.path.join(self.temp_dir, "sans_out.xlsx")
        df, path = convert_pdf_to_excel(pdf_path, output_excel_path=out_xlsx)

        self.assertTrue(os.path.exists(out_xlsx))
        wb = openpyxl.load_workbook(out_xlsx)
        self.assertEqual(wb.sheetnames, ["Sheet1"], "Must be on a single continuous sheet 'Sheet1'")
        ws = wb["Sheet1"]
        self.assertIn("SAUDI AIR NAVIGATION SERVICES", str(ws["C1"].value))
        self.assertIn("INVOICE FOR AIR NAVIGATION CHARGES", str(ws["C2"].value))
        self.assertEqual(len(df), 326, "Must extract all 326 flight records")
        self.assertEqual(len(df.columns), 22, "Must contain all 22 columns")
        self.assertEqual(df.columns[0], "Line Ref #")
        self.assertEqual(df.columns[1], "Item Description")
        self.assertEqual(df.columns[-1], "Item Subtotal (SAR)")

        # Verify financial total of 773,249.47 SAR is present in metadata/totals
        sheet_text = " ".join(str(ws.cell(r, c).value or "") for r in range(1, 20) for c in range(1, 10))
        self.assertIn("773,249.47", sheet_text)

        # Verify Category Summary table at bottom
        last_row_text = str(ws.cell(ws.max_row, 1).value or "")
        self.assertIn("Total Invoice Charges", last_row_text)
        wb.close()

    def test_dubai_airports_invoice(self):
        """Verify specialized conversion for Dubai Airports Corporation invoice (22 pages, 73 flights, 17 cols)."""
        pdf_path = os.path.join(self.samples_dir, "dubai_airports_invoice.pdf")
        if not os.path.exists(pdf_path):
            self.skipTest("dubai_airports_invoice.pdf not found in tests/samples")

        out_xlsx = os.path.join(self.temp_dir, "dubai_out.xlsx")
        df, path = convert_pdf_to_excel(pdf_path, output_excel_path=out_xlsx)

        self.assertTrue(os.path.exists(out_xlsx))
        wb = openpyxl.load_workbook(out_xlsx)
        self.assertEqual(wb.sheetnames, ["Sheet1"], "Must be on a single continuous sheet 'Sheet1'")
        ws = wb["Sheet1"]
        self.assertIn("DUBAI AIRPORTS CORPORATION", str(ws["C1"].value))
        self.assertIn("TAX INVOICE & CUSTOMER WISE AIRCRAFT MOVEMENT REPORT", str(ws["C2"].value))
        self.assertEqual(len(df), 73, "Must extract all 73 flight records")
        self.assertEqual(len(df.columns), 17, "Must contain all 17 columns")
        self.assertEqual(df.columns[0], "Aircraft Type")
        self.assertEqual(df.columns[1], "SL #")
        self.assertEqual(df.columns[-1], "VAT (AED)")

        # Verify summary charges and grand total of 171,251.86 AED
        sheet_text = " ".join(str(ws.cell(r, c).value or "") for r in range(1, 20) for c in range(1, 10))
        self.assertIn("171,251.86", sheet_text)
        self.assertIn("137,170.86", sheet_text)
        self.assertIn("34,081.00", sheet_text)

        # Verify Grand Total Row at bottom of flight table
        grand_total_row_text = " ".join(str(ws.cell(ws.max_row, c).value or "") for c in range(1, 18))
        self.assertIn("GRAND TOTAL", grand_total_row_text)
        self.assertIn("171,251.86", grand_total_row_text)
        wb.close()

    def test_page_without_tables_error_handling(self):
        """Verify document with no tables handles gracefully without crashing."""
        out_xlsx = os.path.join(self.temp_dir, "empty_out.xlsx")
        df, path = convert_pdf_to_excel(self.empty_pdf, output_excel_path=out_xlsx)
        self.assertTrue(os.path.exists(out_xlsx))
        wb = openpyxl.load_workbook(out_xlsx)
        self.assertEqual(wb.sheetnames, ["Sheet1"])
        wb.close()

    def test_storage_janitor_and_purge(self):
        """Verify automatic purge function deletes expired temporary files."""
        # Create a dummy temporary file in OUTPUT_DIR
        dummy_out = os.path.join(OUTPUT_DIR, "test_dummy_expire.xlsx")
        with open(dummy_out, "w") as f:
            f.write("test")

        # Purge with max_age_seconds=0 should immediately delete it
        deleted = purge_expired_files(max_age_seconds=0)
        self.assertFalse(os.path.exists(dummy_out))
        self.assertGreaterEqual(deleted, 1)

    def test_file_not_found(self):
        """Verify FileNotFoundError is raised when input file does not exist."""
        with self.assertRaises(FileNotFoundError):
            convert_pdf_to_excel("non_existent_file.pdf")

    def test_cleaning_helpers(self):
        """Verify cell and table cleaning functions."""
        self.assertEqual(clean_cell_text("  Hello \n World\t "), "Hello World")
        self.assertEqual(clean_cell_text(None), "")
        self.assertEqual(clean_cell_text(1234), "1234")

        dirty_table = [
            ["  Col1 \n ", "Col2\t"],
            ["", None],
            [" Val1 ", "  Val2  "]
        ]
        cleaned = clean_table(dirty_table)
        self.assertEqual(len(cleaned), 2)
        self.assertEqual(cleaned[0], ["Col1", "Col2"])
        self.assertEqual(cleaned[1], ["Val1", "Val2"])


if __name__ == "__main__":
    unittest.main()
